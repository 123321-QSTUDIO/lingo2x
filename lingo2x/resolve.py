"""外部数据解析（@OLE 读写 Excel）与 CALC 段解释执行。

流程位置：parse → resolve_external（本模块）→ analyze/后端。
@OLE 读取需要 openpyxl（只装在 .venv 里，未安装时给出明确报错）。

安全约定：
- @OLE 的文件路径被禁锢在模型文件所在目录内（realpath 校验），
  拒绝 UNC 路径（防 NTLM 中继）与 '..' 逃逸；
- 打开 xlsx 前做 zip 炸弹预检（未压缩总大小 / 压缩比阈值）；
- 写回 Excel 前会把解析后的绝对路径打印出来。
"""
import math
import os
import sys
import zipfile
from itertools import product

from .model import (ModelError, Num, Ref, Bin, Neg, Sum, IVar,
                    QCmp, QAnd, QOr, QNot, CalcAssign, CalcFor, CalcCall,
                    Func, validate_member, MAX_SET_MEMBERS)
from .instantiate import _normkey, _tonum

# zip 炸弹预检阈值
_OLE_MAX_UNCOMPRESSED = 200 * 1024 * 1024   # 未压缩总大小上限 200MB
_OLE_MAX_RATIO = 100.0                       # 单条目压缩比上限


def _safe_ole_path(base_dir, fname):
    """把 @OLE 的文件名解析成绝对路径，并禁锢在 base_dir 内。"""
    if fname.startswith(('\\\\', '//')):
        raise ModelError(f"@OLE 拒绝 UNC 路径（防凭据中继）：{fname!r}")
    base = os.path.realpath(base_dir)
    p = fname if os.path.isabs(fname) else os.path.join(base, fname)
    real = os.path.realpath(p)
    if os.path.commonpath([base, real]) != base:
        raise ModelError(
            f"@OLE 路径越界：{fname!r} 解析为 {real}，不在模型目录 {base} 内。"
            "把数据文件放到模型同目录下再试")
    return real


def _check_zip_bomb(path):
    """打开 xlsx 前的预检：xlsx 是 zip，拒绝解压炸弹。"""
    try:
        with zipfile.ZipFile(path) as zf:
            total = 0
            for info in zf.infolist():
                total += info.file_size
                if info.compress_size > 0 and \
                        info.file_size / info.compress_size > _OLE_MAX_RATIO:
                    raise ModelError(
                        f"@OLE 拒绝可疑的压缩包（疑似 zip 炸弹）："
                        f"{os.path.basename(path)} 条目 {info.filename!r} "
                        f"压缩比 {info.file_size / info.compress_size:.0f}:1")
                if total > _OLE_MAX_UNCOMPRESSED:
                    raise ModelError(
                        f"@OLE 文件解压后过大（>{_OLE_MAX_UNCOMPRESSED // 1024 // 1024}MB），"
                        f"拒绝加载：{os.path.basename(path)}")
    except zipfile.BadZipFile as e:
        raise ModelError(f"@OLE 文件不是有效的 xlsx（zip）：{os.path.basename(path)}（{e}）")


def resolve_external(model, base_dir):
    """依次处理 @OLE 读取（外部数据内联化）和 CALC 段（求解前的参数计算）。"""
    _resolve_ole_reads(model, base_dir)
    _run_calc(model)


# ================= @OLE 读取 =================

def _load_openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        raise ModelError("模型使用了 @OLE 外部数据，需要 openpyxl："
                         ".venv\\Scripts\\python -m pip install openpyxl")


def _member_str(v):
    if isinstance(v, float):
        return f"{v:g}"
    return str(v)


def _range_cells(wb, ref, fname):
    """把区域引用（命名区域 / 'Sheet!A1:B2' / 'A1:B2'）解析成 (工作表, 边界)。"""
    from openpyxl.utils import range_boundaries
    target = None
    for k, dn in wb.defined_names.items():
        if k.lower() == ref.lower():
            target = dn.value
            break
    if target is None:
        # 不是命名区域：必须是合法的单元格/区域引用，否则大概率是用户笔误，
        # 要直接报"找不到命名区域"，而不是把 'CX' 当整列解析出莫名其妙的错误
        import re as _re
        cell_pat = _re.compile(
            r"^(?:'[^']+'|[A-Za-z0-9_一-鿿]+!)?\$?[A-Za-z]{1,3}\$?\d+"
            r"(?::\$?[A-Za-z]{1,3}\$?\d+)?$")
        if not cell_pat.match(ref):
            names = ', '.join(wb.defined_names.keys()) or '（无）'
            raise ModelError(
                f"@OLE 在 {fname} 中找不到命名区域 {ref!r}，它也不是合法的"
                f"单元格区域写法（如 'Sheet1!A1:B2'）。工作簿现有命名区域：{names}")
        target = ref
    if '!' in target:
        sheet, cr = target.split('!', 1)
        sheet = sheet.strip("'")
    else:
        sheet, cr = wb.active.title, target
    cr = cr.replace('$', '')
    if sheet not in wb.sheetnames:
        raise ModelError(f"@OLE 引用 {ref!r} 不存在：{fname} 中没有这个命名区域/工作表")
    if ':' not in cr:                                # 单单元格区域
        cr = f"{cr}:{cr}"
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cr)
    except ValueError as e:
        raise ModelError(f"@OLE 区域引用 {ref!r} 无效：{e}")
    return wb[sheet], (min_col, min_row, max_col, max_row)


def _resolve_ole_reads(model, base_dir):
    if not model.ole_reads:
        return
    openpyxl = _load_openpyxl()
    cache = {}

    def wb_for(fname):
        p = _safe_ole_path(base_dir, fname)
        if not os.path.exists(p):
            raise ModelError(f"@OLE 找不到文件：{p}")
        if p not in cache:
            _check_zip_bomb(p)
            try:
                cache[p] = openpyxl.load_workbook(p, data_only=True)
            except (OSError, ValueError) as e:
                raise ModelError(f"@OLE 无法读取 {os.path.basename(p)}：{e}")
        return cache[p]

    for names, fname, rng in model.ole_reads:
        wb = wb_for(fname)
        for name in names:
            ws, (c0, r0, c1, r1) = _range_cells(wb, rng or name, fname)
            vals = []
            for row in ws.iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1):
                for cell in row:                     # 行主序，与 LINGO 一致
                    if cell.value is None:
                        raise ModelError(f"@OLE 区域 {rng or name} 含空单元格")
                    vals.append(cell.value)
            _assign_external(model, name, vals, fname)


def _assign_external(model, name, vals, fname):
    if name in model.sets:
        sd = model.sets[name]
        if sd.parents:
            raise ModelError(f"派生集 {name} 不能用 @OLE 定义成员")
        if len(vals) > MAX_SET_MEMBERS:
            raise ModelError(f"集合 {name} 从 {fname} 读入的成员数超过上限 "
                             f"{MAX_SET_MEMBERS}（防内存耗尽）")
        mems = [validate_member(_member_str(v), where=f'（{fname} → 集合 {name}）')
                for v in vals]
        if sd.members and sd.members != mems:
            raise ModelError(f"集合 {name} 的成员在 SETS 与 @OLE 中不一致")
        sd.members = mems
        return
    nums = []
    for v in vals:
        if isinstance(v, str):
            try:
                v = float(v)
            except ValueError:
                raise ModelError(f"参数 {name} 从 {fname} 读到了非数字 {v!r}")
        try:
            nums.append(float(v))
        except (TypeError, ValueError):
            raise ModelError(
                f"参数 {name} 从 {fname} 读到了无法转成数字的单元格值 {v!r}"
                f"（类型 {type(v).__name__}）")
    model.data[name] = nums


# ================= @OLE 写回（求解后，scipy 后端调用） =================

def write_ole_back(model, flat, xvector, base_dir):
    """把求解结果写回 Excel。flat 是 FlatLP，xvector 是解向量。"""
    openpyxl = _load_openpyxl()
    pos = {k: i for i, k in enumerate(flat.keys)}
    books = {}
    for fname, rng, names in model.ole_writes:
        p = _safe_ole_path(base_dir, fname)
        if not os.path.exists(p):
            raise ModelError(f"@OLE 写回找不到文件：{p}")
        print(f"@OLE 写回结果到: {p}", file=sys.stderr)
        if p not in books:
            try:
                books[p] = openpyxl.load_workbook(p)
            except (OSError, ValueError) as e:
                raise ModelError(f"@OLE 无法打开写回文件 {os.path.basename(p)}：{e}")
        wb = books[p]
        for name in names:
            dom = _domain_of(model, name)
            vals = []
            for inst in _instances(model, dom):
                key = (name, inst)
                if key not in pos:
                    raise ModelError(f"@OLE 写回的名字 {name} 不是决策变量")
                vals.append(float(xvector[pos[key]]))
            ws, (c0, r0, c1, r1) = _range_cells(wb, rng or name, fname)
            n = (c1 - c0 + 1) * (r1 - r0 + 1)
            if len(vals) != n:
                raise ModelError(f"@OLE 写回区域 {rng or name} 大小（{n}）与 {name} 的实例数（{len(vals)}）不一致")
            it = iter(vals)
            for r in range(r0, r1 + 1):              # 行主序
                for c in range(c0, c1 + 1):
                    ws.cell(row=r, column=c).value = next(it)
    for p, wb in books.items():
        try:
            wb.save(p)
        except OSError as e:
            raise ModelError(f"@OLE 写回保存失败 {os.path.basename(p)}：{e}")


# ================= CALC 段解释执行 =================

def _domain_of(model, name):
    """名字的维度：集合属性 → 父集序列；其他 → ()（标量）。"""
    for sd in model.sets.values():
        if name in sd.attrs:
            return model.set_domain(sd.name)
    return ()


def _instances(model, dom):
    if not dom:
        yield ()
        return
    pools = [model.sets[s].members for s in dom]
    if any(not p for p in pools):
        raise ModelError(f"集合 {'/'.join(dom)} 没有成员")
    yield from product(*pools)


def _eval_idx(e, env):
    if isinstance(e, IVar):
        return env.get(e.name, e.name)
    if isinstance(e, Num):
        return e.value
    if isinstance(e, Neg):
        return -_tonum(_eval_idx(e.x, env))
    if isinstance(e, Bin):
        a = _tonum(_eval_idx(e.l, env))
        b = _tonum(_eval_idx(e.r, env))
        if e.op == '+':
            return a + b
        if e.op == '-':
            return a - b
        if e.op == '*':
            return a * b
        if e.op == '^':
            return a ** b
        if b == 0:
            raise ModelError("CALC 下标表达式除以 0")
        return a / b
    raise ModelError(f"CALC 下标表达式中出现非法节点 {type(e).__name__}")


def _eval_qual(q, env):
    if isinstance(q, QCmp):
        lv, rv = _eval_idx(q.l, env), _eval_idx(q.r, env)
        try:
            lv, rv = _tonum(lv), _tonum(rv)
        except Exception:
            lv, rv = str(lv), str(rv)
            if q.op not in ('=', '!='):
                raise ModelError("CALC 条件过滤中的非数字比较只支持 = / #NE#")
        return {'<=': lv <= rv, '>=': lv >= rv, '=': lv == rv,
                '!=': lv != rv, '<': lv < rv, '>': lv > rv}[q.op]
    if isinstance(q, QAnd):
        return _eval_qual(q.a, env) and _eval_qual(q.b, env)
    if isinstance(q, QOr):
        return _eval_qual(q.a, env) or _eval_qual(q.b, env)
    if isinstance(q, QNot):
        return not _eval_qual(q.a, env)
    raise ModelError("CALC 条件过滤式中出现非法节点")


def _eval_const(model, e, env, values):
    """只求常量值的表达式求值；引用到决策变量或未知名字直接报错。"""
    if isinstance(e, Num):
        return e.value
    if isinstance(e, Ref):
        inst = tuple(_normkey(_eval_idx(x, env)) for x in e.indices)
        key = (e.name, inst)
        if key in values:
            return values[key]
        shown = f"{e.name}({', '.join(inst)})" if inst else e.name
        raise ModelError(f"CALC 段中 {shown} 未定义（CALC 只允许参数参与运算，不能用决策变量）")
    if isinstance(e, Neg):
        return -_eval_const(model, e.x, env, values)
    if isinstance(e, Bin):
        a = _eval_const(model, e.l, env, values)
        b = _eval_const(model, e.r, env, values)
        if e.op == '/':
            if b == 0:
                raise ModelError("CALC 中除以 0")
            return a / b
        if e.op == '^':
            try:
                return a ** b
            except (OverflowError, ZeroDivisionError):
                raise ModelError("CALC 中幂运算溢出或 0 的负次幂")
        return {'+': a + b, '-': a - b, '*': a * b}[e.op]
    if isinstance(e, Sum):
        total = 0.0
        idx = [d[0] for d in e.domain]
        pools = [model.sets[s].members for _, s in e.domain]
        for combo in product(*pools):
            env2 = {**env, **dict(zip(idx, combo))}
            if e.qual is None or _eval_qual(e.qual, env2):
                total += _eval_const(model, e.body, env2, values)
        return total
    if isinstance(e, Func):
        v = _eval_const(model, e.arg, env, values)
        if e.name == 'sqrt':
            if v < 0:
                raise ModelError("CALC 中 @SQRT 的参数为负")
            return math.sqrt(v)
        if e.name == 'abs':
            return abs(v)
        raise ModelError(f"CALC 中暂不支持函数 @{e.name.upper()}")
    raise ModelError(f"CALC 中出现不支持的表达式节点 {type(e).__name__}")


def _exec_block(model, stmts, env, values, active=()):
    for st in stmts:
        if isinstance(st, CalcAssign):
            inst = tuple(_normkey(_eval_idx(x, env)) for x in st.indices)
            values[(st.name, inst)] = _eval_const(model, st.expr, env, values)
            model.calc_assigned.add(st.name)
        elif isinstance(st, CalcCall):
            proc = model.calc_procs.get(st.name)
            if proc is None:
                raise ModelError(f"CALC 调用了未定义的过程 {st.name}")
            if st.name in active:
                raise ModelError(f"CALC 过程 {st.name} 循环调用（{' → '.join(active)}）")
            _exec_block(model, proc, dict(env), values, active + (st.name,))
        elif isinstance(st, CalcFor):
            idx = [d[0] for d in st.domain]
            pools = [model.sets[s].members for _, s in st.domain]
            for combo in product(*pools):
                env2 = {**env, **dict(zip(idx, combo))}
                if st.qual is None or _eval_qual(st.qual, env2):
                    _exec_block(model, st.body, env2, values, active)


def _run_calc(model):
    if not model.calc_main:
        return
    values = {}
    for name, vals in model.data.items():
        dom = _domain_of(model, name)
        for inst, v in zip(_instances(model, dom), vals):
            values[(name, inst)] = float(v)
    _exec_block(model, model.calc_main, {}, values)
    # 写回：CALC 中赋值过的集合属性/标量 → 铺平进 data，供后续 analyze 当作参数
    for name in model.calc_assigned:
        dom = _domain_of(model, name)
        if dom:
            model.data[name] = [values.get((name, inst), 0.0)
                                for inst in _instances(model, dom)]
        else:
            model.data[name] = [values.get((name, ()), 0.0)]
